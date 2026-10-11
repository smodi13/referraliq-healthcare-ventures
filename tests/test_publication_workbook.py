import shutil
import subprocess
import zipfile

import openpyxl
import pytest

from core import config

ORIG = config.WORKBOOK
PUB = config.ROOT / "publication" / "Pilot_ROI_and_Venture_Model_PUBLICATION_COPY.xlsx"
SHEET = "xl/worksheets/sheet10.xml"


def test_only_market_framework_part_differs_byte_for_byte():
    a, b = zipfile.ZipFile(ORIG), zipfile.ZipFile(PUB)
    assert b.testzip() is None and sorted(a.namelist()) == sorted(b.namelist())
    diff = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert diff == [SHEET]
    s = b.read(SHEET).decode()
    assert s.count(',"Not calculated")') == 2 and "PUBLICATION COPY" in s
    assert "D9*D10*D11" in s and "D12*D13" in s  # formulas preserved


def test_original_workbook_still_has_zero_display_formulas():
    s = zipfile.ZipFile(ORIG).read(SHEET).decode()
    assert "D9*D10*D11,0)" in s and "D12*D13,0)" in s


def _recalc(src, tmp_path):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice not available")
    d = tmp_path / "rc"
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(d), str(src)],
                   capture_output=True, timeout=180)
    return openpyxl.load_workbook(d / src.name, data_only=True)


def test_recalculated_publication_copy_matches_original_everywhere_but_the_two_cells(tmp_path):
    o_src = tmp_path / "orig.xlsx"
    shutil.copy(ORIG, o_src)
    p_src = tmp_path / "pub.xlsx"
    shutil.copy(PUB, p_src)
    o, p = _recalc(o_src, tmp_path), _recalc(p_src, tmp_path)
    assert o.sheetnames == p.sheetnames
    diffs = []
    for name in o.sheetnames:
        a, b = o[name], p[name]
        for row in a.iter_rows():
            for c in row:
                v2 = b[c.coordinate].value
                if c.value != v2:
                    diffs.append((name, c.coordinate, c.value, v2))
    assert {(s, c) for s, c, _, _ in diffs} == {("MARKET FRAMEWORK", "D12"), ("MARKET FRAMEWORK", "D14")}
    assert p["MARKET FRAMEWORK"]["D12"].value == "Not calculated" and o["MARKET FRAMEWORK"]["D12"].value == 0


def test_formula_is_live_once_inputs_are_supplied(tmp_path):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice not available")
    wb = openpyxl.load_workbook(PUB)  # temp copy only; formulas preserved
    ws = wb["MARKET FRAMEWORK"]
    ws["D9"], ws["D10"], ws["D11"], ws["D13"] = 100, 0.5, 50000, 0.1
    src = tmp_path / "filled.xlsx"
    wb.save(src)
    out = _recalc(src, tmp_path)["MARKET FRAMEWORK"]
    assert out["D12"].value == 100 * 0.5 * 50000 and out["D14"].value == pytest.approx(100 * 0.5 * 50000 * 0.1)
